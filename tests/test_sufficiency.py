"""证据闸门：能说出理由才拒答，且拒答真的会挡住大模型。

这一层最容易被写成"看起来接上了"：判据算出来了、meta 里也有，候选照样被送进
大模型；或者反过来，判据太宽，把本来答得出的题一起拒掉。所以断言落在两件事上：

1. `judge` 的阻断/放行由**哪一条**信号决定（逐条单独验，不混在一起）；
2. `stream_answer` 在 `mode=on` 下**一次 LLM 都没调**、`mode=warn` 下答案照常。

第 2 条是这层的全部价值：只写进 meta 的拒答等于没拒。
"""

from __future__ import annotations

import sqlite3
from collections.abc import AsyncIterator
from pathlib import Path

from rag4prefa import ask, sufficiency
from rag4prefa.adapters.vectors import NumpyStore
from rag4prefa.config import Settings
from rag4prefa.domain import Filter, Hit
from rag4prefa.ingest import ingest_file
from tests.conftest import FakeEmbedder

QUESTION = "张拉控制要点有哪些？"
OUT_OF_SCOPE = "北京市装配化装修的补贴标准是多少？"


def _hit(chunk_id: str = "c1", *, rerank: float | None = None, dense: float | None = None) -> Hit:
    return Hit(chunk_id=chunk_id, text="正文", rerank_score=rerank, dense_score=dense)


def _mentions(known: dict[str, int]):
    return lambda needle: known.get(needle, 0)


# ── 判据逐条 ────────────────────────────────────────────────


def test_no_hits_refuses(conn) -> None:
    verdict = sufficiency.judge(conn, QUESTION, [])
    assert verdict.refuse is True
    assert [s.name for s in verdict.signals] == ["no_hits"]


def test_place_absent_from_corpus_refuses(conn) -> None:
    """库里一个字都没提过这个地域 → 拒。

    这是库外集里 4/6 道题的形态（北京问出宁波的方案、日本问出绍兴的图集）。
    判据用"全库提没提过"，不用"候选像不像"：向量检索永远会返回最近的一条。
    """
    verdict = sufficiency.judge(
        conn, OUT_OF_SCOPE, [_hit()], corpus_mentions=_mentions({"北京": 0})
    )
    assert verdict.refuse is True
    assert [s.name for s in verdict.signals] == ["geo_absent"]
    assert "北京" in verdict.reasons[0]


def test_place_present_in_corpus_does_not_refuse(conn) -> None:
    """同一个问句，库里提过北京 → 不因地域拒（至于答不答得对，是另一层的事）。"""
    verdict = sufficiency.judge(
        conn, OUT_OF_SCOPE, [_hit()], corpus_mentions=_mentions({"北京": 12})
    )
    assert verdict.refuse is False
    assert verdict.signals == []


def test_only_the_absent_place_is_named(conn) -> None:
    """问句里出现多个地域时，只点名库里没有的那个。"""
    verdict = sufficiency.judge(
        conn,
        "浙江和江苏的补贴标准有什么不同？",
        [_hit()],
        corpus_mentions=_mentions({"浙江": 900, "江苏": 0}),
    )
    assert verdict.refuse is True
    detail = verdict.reasons[0]
    assert "江苏" in detail
    assert "浙江" not in detail


def test_low_top_score_refuses(conn) -> None:
    verdict = sufficiency.judge(
        conn, QUESTION, [_hit(rerank=0.01)], min_rerank=0.05
    )
    assert verdict.refuse is True
    assert [s.name for s in verdict.signals] == ["weak_top"]


def test_top_score_above_floor_passes(conn) -> None:
    verdict = sufficiency.judge(conn, QUESTION, [_hit(rerank=0.30)], min_rerank=0.05)
    assert verdict.refuse is False


def test_two_thresholds_are_not_interchangeable(conn) -> None:
    """重排分与余弦相似度量纲不同，门槛必须各管各的。

    合成一个门槛时，最典型的事故是"重排开着时用余弦门槛"：余弦 0.62 看着很高，
    而重排分 0.02 才是同一批候选的真实判断——用错门槛会让闸门彻底失效。
    """
    # 有重排分时，只看重排门槛：0.9 的余弦救不了 0.01 的重排分
    verdict = sufficiency.judge(
        conn, QUESTION, [_hit(rerank=0.01, dense=0.90)], min_rerank=0.05, min_dense=0.5
    )
    assert verdict.refuse is True

    # 没有重排分时才退回余弦门槛；此时重排门槛不参与
    verdict = sufficiency.judge(conn, QUESTION, [_hit(dense=0.60)], min_rerank=0.99, min_dense=0.5)
    assert verdict.refuse is False


def test_threshold_zero_means_no_score_gate(conn) -> None:
    """门槛 0 是"不设分闸"，不是"必须大于 0"。否则关掉闸门要靠改代码。"""
    verdict = sufficiency.judge(conn, QUESTION, [_hit(rerank=0.0)], min_rerank=0.0)
    assert verdict.refuse is False


def test_explain_says_what_is_missing(conn) -> None:
    verdict = sufficiency.judge(
        conn, OUT_OF_SCOPE, [_hit()], corpus_mentions=_mentions({"北京": 0})
    )
    text = sufficiency.explain(verdict)
    assert "资料不足" in text
    assert "北京" in text
    fine = sufficiency.judge(conn, QUESTION, [_hit()], corpus_mentions=_mentions({"北京": 1}))
    assert sufficiency.explain(fine) == ""


# ── 生产链路：拒答必须真的挡住大模型 ─────────────────────────


class _CountingLLM:
    def __init__(self) -> None:
        self.calls = 0

    async def stream(self, system: str, prompt: str) -> AsyncIterator[str]:
        self.calls += 1
        yield "按资料，张拉控制要点如下。"


def _seed(conn: sqlite3.Connection, settings: Settings, docx: Path) -> NumpyStore:
    store = NumpyStore(conn, FakeEmbedder.dim)
    outcome = ingest_file(
        conn, store, FakeEmbedder(), settings, path=docx, project_name="甲项目", doc_year=2024
    )
    assert outcome.status == "indexed", outcome.error
    return store


def _events(conn, store, settings, question: str, llm) -> list[tuple[str, dict]]:
    async def run():
        return [
            item
            async for item in ask.stream_answer(
                conn,
                store,
                FakeEmbedder(),
                llm,
                None,  # 不重排：这条用例只验闸门
                settings,
                question=question,
                filt=Filter(),
                top_k=3,
            )
        ]

    import asyncio

    return asyncio.run(run())


def test_stream_answer_refuses_without_calling_the_model(
    conn, settings: Settings, docx_file: Path
) -> None:
    store = _seed(conn, settings, docx_file)
    llm = _CountingLLM()
    events = _events(conn, store, settings, OUT_OF_SCOPE, llm)

    assert llm.calls == 0
    names = [name for name, _ in events]
    assert "citations" in names
    citations = dict(events)["citations"]
    assert citations == []
    text = "".join(p["text"] for n, p in events if n == "token")
    assert "资料不足" in text
    meta = dict(events)["meta"]
    assert meta["refusal"]["refuse"] is True
    assert meta["refuse_mode"] == "on"


def test_warn_mode_records_but_answers(conn, settings: Settings, docx_file: Path) -> None:
    """`warn` 是量误拒率用的形态：判断照算、写进 meta，但答案不受影响。"""
    store = _seed(conn, settings, docx_file)
    llm = _CountingLLM()
    warn = settings.model_copy(update={"refuse_mode": "warn"})
    events = _events(conn, store, warn, OUT_OF_SCOPE, llm)

    assert llm.calls == 1
    assert dict(events)["meta"]["refusal"]["refuse"] is True
    assert dict(events)["citations"] != []


def test_off_mode_does_not_even_judge(conn, settings: Settings, docx_file: Path) -> None:
    """`off` 是"与加这一层之前完全一致"，连 meta 里都不该出现 refusal 判断。"""
    store = _seed(conn, settings, docx_file)
    llm = _CountingLLM()
    off = settings.model_copy(update={"refuse_mode": "off"})
    events = _events(conn, store, off, OUT_OF_SCOPE, llm)

    assert llm.calls == 1
    assert dict(events)["meta"]["refusal"] is None


def test_in_scope_question_still_answers(conn, settings: Settings, docx_file: Path) -> None:
    """误拒是这一层的主要风险，用库里真有的题钉住"不该拒的不拒"。"""
    store = _seed(conn, settings, docx_file)
    llm = _CountingLLM()
    events = _events(conn, store, settings, QUESTION, llm)

    assert llm.calls == 1
    assert dict(events)["meta"]["refusal"]["refuse"] is False


def test_place_in_body_is_not_coverage(conn, settings: Settings, tmp_path: Path) -> None:
    """正文提到 ≠ 库里有这个地域的资料。这是这套判据最容易做错的一处。

    实际语料里"北京"出现在 17 个片段中（参编单位名单、企业简介、区域规划里的接轨），
    但没有一份文件是关于北京的。拿正文判断，库外问题会被判成"库里有"，闸门永远不
    触发——而它不会报错，只会安静地全部放行。这里用一份"正文提到上海、文件名与上海
    无关"的文件把这两件事分开。
    """
    from docx import Document

    from rag4prefa import repo

    doc = Document()
    doc.add_heading("第1章 企业简介", level=1)
    doc.add_paragraph("公司承建了上海某会展中心、北京某峰会等标志性工程。" * 5)
    doc.add_heading("第2章 技术要点", level=2)
    doc.add_paragraph("本项目采用装配式内装修，设备管线与主体结构分离设置。" * 5)
    path = tmp_path / "企业简介.docx"
    doc.save(path)

    store = NumpyStore(conn, FakeEmbedder.dim)
    outcome = ingest_file(
        conn, store, FakeEmbedder(), settings, path=path, project_name="甲项目", doc_year=2024
    )
    assert outcome.status == "indexed", outcome.error

    assert repo.chunks_mentioning(conn, "上海") > 0  # 正文里有
    assert repo.documents_mentioning(conn, "上海") == 0  # 但没有一份文件是关于它的
    assert repo.documents_mentioning(conn, "企业简介") > 0  # 文件名命中才算
    assert repo.documents_mentioning(conn, "") == 0
    assert repo.chunks_mentioning(conn, "") == 0

    # 不走注入：直接看闸门自己调的是哪个口径。改成 `chunks_mentioning` 这条会红。
    verdict = sufficiency.judge(conn, "北京市的装配化装修补贴标准是多少？", [_hit()])
    assert verdict.refuse is True, "正文提到上海/北京不等于库里有这个地域的资料"


def test_geo_gate_uses_filename_coverage_not_body(conn) -> None:
    """闸门必须用"有没有关于它的文件"，不能用"正文提没提过"。"""
    verdict = sufficiency.judge(
        conn,
        OUT_OF_SCOPE,
        [_hit()],
        corpus_mentions=_mentions({"北京": 0}),  # 文件名口径：没有北京的文件
    )
    assert verdict.refuse is True


# ── 地域层级：省级文件覆盖省内城市 ──────────────────────────

HANGZHOU = "杭州的装配式内装修施工有什么要求？"


def test_city_covered_by_province_file_does_not_refuse(conn) -> None:
    """库里只有浙江省的地方标准、文件名里没有"杭州"——省级标准适用于省内城市，不能拒。

    这是专家复审第四节点的误拒：拿"文件名必须出现城市名"当判据，会把"有省标可查"
    的问题判成"库里没有杭州的资料"。误拒比漏拒更贵。
    """
    verdict = sufficiency.judge(
        conn, HANGZHOU, [_hit()],
        corpus_mentions=_mentions({"杭州": 0, "浙江": 22}),
    )
    assert verdict.refuse is False
    # 不拒，但要把"这是省级口径"说出来：答案能被读成杭州市自己的规定才是真问题。
    names = {s.name: s for s in verdict.signals}
    assert set(names) == {"geo_upper_level"}
    assert names["geo_upper_level"].blocking is False
    assert "浙江" in names["geo_upper_level"].detail


def test_city_with_own_files_has_no_geo_signal(conn) -> None:
    verdict = sufficiency.judge(
        conn, HANGZHOU, [_hit()], corpus_mentions=_mentions({"杭州": 3, "浙江": 22})
    )
    assert verdict.refuse is False
    assert verdict.signals == []


def test_city_without_province_files_still_refuses(conn) -> None:
    """省里也没有这个市的资料 → 照旧拒。放宽只允许"上一级真的有文件"这一种情形。"""
    verdict = sufficiency.judge(
        conn, "成都的装配式装修验收规定是什么？", [_hit()],
        corpus_mentions=_mentions({"成都": 0, "四川": 0}),
    )
    assert verdict.refuse is True
    assert [s.name for s in verdict.signals] == ["geo_absent"]


def test_municipality_has_no_upper_level_coverage(conn) -> None:
    """直辖市不挂在"全国"上：库里有国家标准，不等于有北京的申报口径。

    这条是 7/7 地域拒答的护栏——把直辖市挂到全国，"北京市的补贴标准"就会因为库里有
    GB 标准而放行，那一类拒答会全塌。
    """
    verdict = sufficiency.judge(
        conn, "北京市装配化装修的补贴标准是多少？", [_hit()],
        corpus_mentions=_mentions({"北京": 0, "全国": 40, "国家": 40}),
    )
    assert verdict.refuse is True
    assert [s.name for s in verdict.signals] == ["geo_absent"]
    assert sufficiency.coverage_place("北京") is None


def test_place_hierarchy_covers_the_cities_actually_in_the_corpus() -> None:
    """表要对得上语料的实际覆盖面，否则"省覆盖市"只是纸上的一句话。

    只写"杭州归浙江"这一条，把 `PROVINCE_OF` 删到只剩一条也能全绿；这里对着
    `PLACES` 里的浙江城市逐个数。
    """
    for city in ("杭州", "宁波", "温州", "绍兴", "湖州", "嘉兴", "金华", "台州",
                 "丽水", "衢州", "舟山"):
        assert sufficiency.coverage_place(city) == "浙江", city
    # 省级自己不需要上级；表里写 `"浙江": "浙江"` 只是为了自解释，取值时必须当成"无上级"。
    assert sufficiency.coverage_place("浙江") is None
    # 境外体系没有上级可言。
    assert sufficiency.coverage_place("日本") is None
