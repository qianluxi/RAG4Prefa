from __future__ import annotations

import sqlite3
from pathlib import Path

import pytest

from rag4prefa import repo
from rag4prefa.adapters.vectors import NumpyStore
from rag4prefa.config import Settings
from rag4prefa.ingest import ingest_file
from tests.conftest import FakeEmbedder


def _store(conn: sqlite3.Connection) -> NumpyStore:
    return NumpyStore(conn, FakeEmbedder.dim)


def test_ingest_builds_index_and_facts(
    conn: sqlite3.Connection, settings: Settings, docx_file: Path
) -> None:
    embedder = FakeEmbedder()
    store = _store(conn)
    outcome = ingest_file(
        conn, store, embedder, settings, path=docx_file, project_name="测试项目", doc_year=2024
    )

    assert outcome.status == "indexed"
    assert outcome.chunks > 0
    assert store.count() == outcome.chunks

    project = repo.get_project(conn, "测试项目")
    assert project is not None
    assert project.document_count == 1
    assert project.chunk_count == outcome.chunks  # 统计来自 COUNT(*)，不再是硬编码 0


def test_reingest_same_file_is_skipped(
    conn: sqlite3.Connection, settings: Settings, docx_file: Path
) -> None:
    embedder = FakeEmbedder()
    store = _store(conn)
    first = ingest_file(
        conn, store, embedder, settings, path=docx_file, project_name="测试项目", doc_year=2024
    )
    second = ingest_file(
        conn, store, embedder, settings, path=docx_file, project_name="测试项目", doc_year=2024
    )

    assert first.status == "indexed"
    assert second.status == "skipped"
    assert store.count() == first.chunks  # 没有重复向量


def test_force_reingest_does_not_duplicate(
    conn: sqlite3.Connection, settings: Settings, docx_file: Path
) -> None:
    embedder = FakeEmbedder()
    store = _store(conn)
    first = ingest_file(
        conn, store, embedder, settings, path=docx_file, project_name="测试项目", doc_year=2024
    )
    ingest_file(
        conn,
        store,
        embedder,
        settings,
        path=docx_file,
        project_name="测试项目",
        doc_year=2024,
        force=True,
    )
    assert store.count() == first.chunks


def test_delete_project_removes_vectors_and_rows(
    conn: sqlite3.Connection, settings: Settings, docx_file: Path
) -> None:
    embedder = FakeEmbedder()
    store = _store(conn)
    ingest_file(
        conn, store, embedder, settings, path=docx_file, project_name="测试项目", doc_year=2024
    )
    chunk_ids = repo.delete_project(conn, "测试项目")
    store.delete(chunk_ids)

    assert repo.get_project(conn, "测试项目") is None
    assert store.count() == 0
    assert repo.fts_search(conn, "悬挑", limit=5) == []


def test_search_filters_by_project(
    conn: sqlite3.Connection, settings: Settings, docx_file: Path
) -> None:
    embedder = FakeEmbedder()
    store = _store(conn)
    ingest_file(
        conn, store, embedder, settings, path=docx_file, project_name="甲项目", doc_year=2024
    )
    ingest_file(
        conn, store, embedder, settings, path=docx_file, project_name="乙项目", doc_year=2020
    )

    from rag4prefa.domain import Filter

    all_hits = store.search(embedder.embed_one("悬挑"), limit=50)
    only_jia = store.search(embedder.embed_one("悬挑"), limit=50, filt=Filter(projects=["甲项目"]))
    assert len(only_jia) < len(all_hits)
    assert len(only_jia) > 0

    docs = repo.get_chunks(conn, [cid for cid, _ in only_jia])
    assert {d["project_name"] for d in docs.values()} == {"甲项目"}


def test_fts_search_finds_chinese_terms(
    conn: sqlite3.Connection, settings: Settings, docx_file: Path
) -> None:
    embedder = FakeEmbedder()
    store = _store(conn)
    ingest_file(
        conn, store, embedder, settings, path=docx_file, project_name="测试项目", doc_year=2024
    )
    hits = repo.fts_search(conn, "缓粘结预应力", limit=10)
    assert hits, "BM25 应该能命中中文专有名词"


def test_index_can_be_rebuilt_from_sqlite(
    conn: sqlite3.Connection, settings: Settings, docx_file: Path
) -> None:
    from rag4prefa.ingest import reindex

    embedder = FakeEmbedder()
    store = _store(conn)
    ingest_file(
        conn, store, embedder, settings, path=docx_file, project_name="测试项目", doc_year=2024
    )
    before = store.count()

    store.clear()
    assert store.count() == 0

    result = reindex(conn, store, embedder, settings)
    assert result["chunks"] == before
    assert store.count() == before


class FailingEmbedder(FakeEmbedder):
    """模拟限流 / 超时 / 密钥失效：远程调用一定会抛。"""

    def embed(self, texts: list[str]) -> list[list[float]]:
        raise RuntimeError("429 rate limited")


def _update_docx(path: Path, extra: str) -> None:
    from docx import Document

    doc = Document(path)
    doc.add_paragraph(extra)
    doc.save(path)


def test_embedding_failure_keeps_previous_version_searchable(
    conn: sqlite3.Connection, settings: Settings, docx_file: Path
) -> None:
    """更新文档时 embedding 失败，旧版本必须还能被检索到。

    旧实现先删旧片段再调 embedding，一次限流就把原先能用的版本弄没了。
    """
    store = _store(conn)
    first = ingest_file(
        conn, store, FakeEmbedder(), settings, path=docx_file, project_name="测试项目"
    )
    assert first.status == "indexed"
    before_vectors = store.count()
    before_chunks = repo.chunk_count(conn)
    assert before_vectors == before_chunks > 0

    _update_docx(docx_file, "新增一段关于屋面防水的补充说明。")
    failed = ingest_file(
        conn, store, FailingEmbedder(), settings, path=docx_file, project_name="测试项目"
    )

    assert failed.status == "failed"
    assert "rate limited" in (failed.error or "")
    assert store.count() == before_vectors, "失败不该删掉旧向量"
    assert repo.chunk_count(conn) == before_chunks, "失败不该删掉旧片段"
    assert repo.fts_search(conn, "缓粘结预应力", limit=5), "旧版本仍应可检索"
    assert repo.get_document(conn, "测试项目", "方案.docx").chunk_count == before_chunks


def test_retry_after_failure_is_not_skipped(
    conn: sqlite3.Connection, settings: Settings, docx_file: Path
) -> None:
    """失败过再重试必须真的重跑，而不是因为 content_sha 相同被跳过。"""
    store = _store(conn)
    ingest_file(conn, store, FakeEmbedder(), settings, path=docx_file, project_name="测试项目")
    _update_docx(docx_file, "补充：陶土砖反打工艺的验收要点。")
    failed = ingest_file(
        conn, store, FailingEmbedder(), settings, path=docx_file, project_name="测试项目"
    )
    assert failed.status == "failed"

    retried = ingest_file(
        conn, store, FakeEmbedder(), settings, path=docx_file, project_name="测试项目"
    )

    assert retried.status == "indexed", "失败后的重试被误判成 skipped"
    assert repo.get_document(conn, "测试项目", "方案.docx").status == "indexed"
    assert repo.fts_search(conn, "陶土砖反打", limit=5), "新版本内容应已入库"


def test_reingest_smaller_version_drops_stale_chunks(
    conn: sqlite3.Connection, settings: Settings, tmp_path: Path
) -> None:
    """新版本片段更少时，多出来的旧片段与旧向量都要被清掉。"""
    from docx import Document

    long_doc = tmp_path / "长文.docx"
    doc = Document()
    doc.add_heading("第1章 结构", level=1)
    for i in range(6):
        doc.add_paragraph(f"第{i}段：悬挑构件的预应力张拉控制要点与摩阻损失。" * 4)
    doc.save(long_doc)

    store = _store(conn)
    big = ingest_file(
        conn, store, FakeEmbedder(), settings, path=long_doc, project_name="测试项目"
    )
    assert big.status == "indexed"
    assert big.chunks > 1

    short_doc = tmp_path / "长文.docx"
    doc = Document()
    doc.add_heading("第1章 结构", level=1)
    doc.add_paragraph("只剩一段：悬挑构件的预应力张拉控制要点。")
    doc.save(short_doc)

    small = ingest_file(
        conn, store, FakeEmbedder(), settings, path=short_doc, project_name="测试项目"
    )

    assert small.status == "indexed"
    assert small.chunks < big.chunks
    assert store.count() == small.chunks, "多出来的旧向量必须被删除"
    assert repo.chunk_count(conn) == small.chunks
    # 索引里剩下的每个 id 都还能在事实表里找到正文，不留指向已删片段的历史向量
    hits = store.search(FakeEmbedder().embed_one("悬挑"), limit=50)
    assert {cid for cid, _ in hits} == set(repo.get_chunks(conn, [c for c, _ in hits]))


# ── 故障注入：切换过程中每一步失败，用户还剩下什么 ──────────────────


class HalfWritingStore(NumpyStore):
    """写到一半就抛：模拟向量索引写入中途失败。"""

    def upsert(self, items) -> None:  # type: ignore[override]
        super().upsert(list(items)[: max(1, len(items) // 2)])
        raise RuntimeError("模拟索引写入中途失败")


class NoDeleteStore:
    """独立存储的假后端（像 Milvus 那样自己存一份 id）：删旧向量时抛错。

    刻意**不继承** `NumpyStore`：numpy 的索引就是 `vectors` 表的视图，而切换事实时
    表里的旧行已经被同一个事务删掉了，它表达不出"索引比归档多"这个状态——
    而那正是本用例要测的东西（只有独立存储的后端才会留下这种残留）。
    """

    name = "no-delete"
    notes: list[str] = []

    def __init__(self, conn, dim: int) -> None:
        self.dim = dim
        self._ids: set[str] = set()

    def upsert(self, items) -> None:  # type: ignore[override]
        self._ids |= {i.chunk_id for i in items}

    def delete(self, chunk_ids) -> None:  # type: ignore[override]
        raise RuntimeError("模拟删除旧向量失败")

    def search(self, vector, *, limit, filt=None):  # type: ignore[override]
        return []

    def count(self) -> int:
        return len(self._ids)

    def ids(self) -> set[str]:
        return set(self._ids)

    def clear(self) -> None:
        self._ids = set()

    def rebuild(self, items) -> None:  # type: ignore[override]
        self._ids = {i.chunk_id for i in items}


def _state_is_clean(conn: sqlite3.Connection, store: NumpyStore) -> bool:
    from rag4prefa.consistency import inspect

    return inspect(conn, store_count=store.count()).blockers == []


def test_failure_before_fact_switch_keeps_old_version_intact(
    conn: sqlite3.Connection, settings: Settings, docx_file: Path, monkeypatch
) -> None:
    """事实切换失败时：旧版本必须原样可检索，且不能留下"没有正文的向量"。

    这是"先写索引、再切事实"顺序的价值所在——写索引失败时什么都还没变；
    切事实失败时把多写的向量撤回，库回到上一版，而不是停在"索引多了几条"。
    """
    store = _store(conn)
    first = ingest_file(
        conn, store, FakeEmbedder(), settings, path=docx_file, project_name="测试项目"
    )
    assert first.status == "indexed"
    before_vectors = store.count()
    before_chunks = repo.chunk_count(conn)

    _update_docx(docx_file, "新增一段关于屋面防水的补充说明。")

    def boom(*args, **kwargs):
        raise RuntimeError("模拟切换片段时数据库出错")

    monkeypatch.setattr(repo, "commit_document_version", boom)
    failed = ingest_file(
        conn, store, FakeEmbedder(), settings, path=docx_file, project_name="测试项目"
    )

    assert failed.status == "failed"
    assert "已撤回本次新增的" in (failed.error or "")
    assert repo.chunk_count(conn) == before_chunks, "旧片段不能被删"
    assert store.count() == before_vectors, "多写的向量必须撤回"
    assert _state_is_clean(conn, store), "撤回之后索引必须仍然自洽，否则检索会被无谓地阻断"
    assert repo.fts_search(conn, "缓粘结预应力", limit=5), "旧版本仍应可检索"


class FullWriteThenRaiseStore(NumpyStore):
    """全部写完之后才抛：模拟"最后一批写完、连接断在返回路上"。

    这是最狠的部分写入——残留里**一定**包含本次新增的 chunk_id，正好用来验证
    撤回是不是按"写入前算好的清单"执行。只写前半批的那种故障，前半批很可能全是
    沿用旧 id 的片段（内容没变的段落），撤回与不撤回看起来一样，测不出问题。
    """

    def upsert(self, items) -> None:  # type: ignore[override]
        super().upsert(list(items))
        raise RuntimeError("模拟最后一批写入后连接断开")


def test_partial_write_rollback_covers_newly_created_chunks(
    conn: sqlite3.Connection, settings: Settings, docx_file: Path
) -> None:
    """故障发生前确实写入了新 id，撤回也必须把它删干净。

    为什么这条要单独测：`written` 曾经是在 `store.upsert()` 正常返回之后才算的，
    而 upsert 抛异常时根本不会返回——撤回清单永远是空的，逻辑等于不存在。残留的
    向量条数多于归档，检索会被一致性检查拒绝：不会给出错答案，但"旧版本照常可用"
    这个承诺没兑现。
    """
    store = _store(conn)
    first = ingest_file(
        conn, store, FakeEmbedder(), settings, path=docx_file, project_name="测试项目"
    )
    assert first.status == "indexed"
    before_chunks = repo.chunk_count(conn)
    before_ids = repo.archived_chunk_ids(conn)

    _update_docx(docx_file, "补充：陶土砖反打工艺的验收要点。")

    broken = FullWriteThenRaiseStore(conn, FakeEmbedder.dim)
    failed = ingest_file(
        conn, broken, FakeEmbedder(), settings, path=docx_file, project_name="测试项目"
    )

    assert failed.status == "failed"
    assert "已撤回本次新增的" in (failed.error or "")
    # 这条用例的前提：故障发生前确实写入了新 id。否则它和"什么都没写"的用例等价，
    # 撤回清单是不是空的就都看不出来了。
    rolled_back = int((failed.error or "").split("已撤回本次新增的 ")[1].split(" 条")[0])
    assert rolled_back > 0, "这份文档的修改必须真的产生新 chunk_id"
    # 逐 id 比对，而不是只比条数：新 id 进去了、旧 id 掉了一个，条数照样可能相等
    assert repo.archived_chunk_ids(conn) == before_ids
    assert repo.chunk_count(conn) == before_chunks, "事实表不该有任何变化"
    assert repo.get_index_state(conn) == repo.INDEX_READY, "入库不该动索引切换标记"
    clean = _store(conn)
    assert _state_is_clean(conn, clean), "撤回之后索引必须仍然自洽"
    assert repo.fts_search(conn, "缓粘结预应力", limit=5), "旧版本仍应可检索"


def test_index_failure_midway_leaves_previous_version_searchable(
    conn: sqlite3.Connection, settings: Settings, docx_file: Path
) -> None:
    """索引写到一半就失败：事实不该被改动，旧版本必须完整可用。

    注意它覆盖的是"只写完前半批"的情形：这份文档是在末尾追加段落，前半批大多是
    内容未变、沿用旧 id 的片段，所以**不能**用它证明"新 id 被撤回了"——
    那件事由 `test_partial_write_rollback_covers_newly_created_chunks` 负责。
    """
    first_store = _store(conn)
    first = ingest_file(
        conn, first_store, FakeEmbedder(), settings, path=docx_file, project_name="测试项目"
    )
    assert first.status == "indexed"
    before_chunks = repo.chunk_count(conn)
    before_vectors = repo.vector_stats(conn)["count"]

    _update_docx(docx_file, "补充：陶土砖反打工艺的验收要点。")
    broken = HalfWritingStore(conn, FakeEmbedder.dim)
    broken.dim = FakeEmbedder.dim
    failed = ingest_file(
        conn, broken, FakeEmbedder(), settings, path=docx_file, project_name="测试项目"
    )

    assert failed.status == "failed"
    assert repo.chunk_count(conn) == before_chunks, "事实表不该有任何变化"
    assert repo.vector_stats(conn)["count"] == before_vectors, "写入一半的向量必须被撤回"
    clean = _store(conn)
    assert _state_is_clean(conn, clean), "残留会让检索被拒绝，必须清干净"
    assert repo.fts_search(conn, "缓粘结预应力", limit=5)


def test_stale_vector_cleanup_failure_is_not_reported_as_success(
    conn: sqlite3.Connection, settings: Settings, tmp_path: Path
) -> None:
    """删旧向量失败：不能把这次切换标成成功，且检索要被拒绝（索引比归档多）。"""
    from docx import Document

    from rag4prefa import consistency
    from rag4prefa.consistency import IndexMismatch

    long_doc = tmp_path / "长文.docx"
    doc = Document()
    doc.add_heading("第1章 结构", level=1)
    for i in range(6):
        doc.add_paragraph(f"第{i}段：悬挑构件的预应力张拉控制要点与摩阻损失。" * 4)
    doc.save(long_doc)

    store = NoDeleteStore(conn, FakeEmbedder.dim)
    big = ingest_file(
        conn, store, FakeEmbedder(), settings, path=long_doc, project_name="测试项目"
    )
    assert big.status == "indexed"

    short_doc = tmp_path / "长文.docx"
    doc = Document()
    doc.add_heading("第1章 结构", level=1)
    doc.add_paragraph("只剩一段：悬挑构件的预应力张拉控制要点。")
    doc.save(short_doc)

    failed = ingest_file(
        conn, store, FakeEmbedder(), settings, path=short_doc, project_name="测试项目"
    )

    assert failed.status == "failed", "切换没干净收尾，不能谎报成功"
    assert "新版本已生效" in (failed.error or "")
    assert repo.get_document(conn, "测试项目", "长文.docx").status == "failed"
    # 新内容确实写进去了，但旧向量没删掉 → 索引比归档多，检索必须被拒绝
    with pytest.raises(IndexMismatch, match="向量索引实际"):
        consistency.ensure_queryable(
            conn, FakeEmbedder.model, store.dim, store_count=store.count()
        )
